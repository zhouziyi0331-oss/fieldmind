"""
PowerPoint 采集插件
"""

from typing import Dict, Any
from app.agents.ingestion_agent import IngestionPlugin
from app.core.logging import logger


class PPTPlugin(IngestionPlugin):
    """PowerPoint 采集插件"""

    @property
    def plugin_name(self) -> str:
        return "PPTPlugin"

    @property
    def supported_formats(self) -> list:
        return ["ppt", "pptx"]

    def ingest(self, file_path: str, metadata: Dict[str, Any]) -> Dict[str, Any]:
        """采集 PowerPoint 文件"""
        try:
            from pptx import Presentation

            prs = Presentation(file_path)

            # 提取元数据
            structured_metadata = {
                "title": prs.core_properties.title or "",
                "author": prs.core_properties.author or "",
                "subject": prs.core_properties.subject or "",
                "slides": len(prs.slides),
            }

            # 提取文本
            text_content = []

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []

                # 提取标题
                if slide.shapes.title:
                    slide_text.append(f"【幻灯片 {slide_num} 标题】")
                    slide_text.append(slide.shapes.title.text)

                # 提取所有文本框
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)

                if slide_text:
                    text_content.append("\n".join(slide_text))

            full_text = "\n\n".join(text_content)

            # 统计
            structured_metadata["total_words"] = len(full_text.replace(" ", ""))
            structured_metadata["total_sentences"] = full_text.count('。') + full_text.count('.')
            structured_metadata["language"] = self._detect_language(full_text)

            return {
                "raw_text": full_text,
                "structured_metadata": structured_metadata,
                "content_type": "document",
                "extraction_method": "python-pptx",
                "confidence": 0.9
            }

        except Exception as e:
            logger.error(f"PPT 采集失败: {e}")
            raise

    def _detect_language(self, text: str) -> str:
        chinese_chars = sum(1 for c in text if '一' <= c <= '鿿')
        total_chars = len(text)
        if total_chars > 0 and chinese_chars / total_chars > 0.3:
            return "zh"
        return "en"
