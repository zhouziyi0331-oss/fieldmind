"""
文档内容提取器
支持多种文档格式的文本提取
"""

import os
from typing import Optional, Dict, Any, List
from pathlib import Path
from io import BytesIO

from app.core.logging import logger
from app.core.exceptions import DocumentProcessingException
from app.core.errors import ErrorCode


class DocumentExtractor:
    """文档内容提取器"""

    @staticmethod
    def extract(
        file_path: str,
        mime_type: str
    ) -> Dict[str, Any]:
        """
        提取文档内容（统一接口）

        Args:
            file_path: 文件路径
            mime_type: MIME类型

        Returns:
            Dict: 提取结果
                {
                    "text": str,              # 提取的文本
                    "metadata": dict,         # 元数据
                    "pages": int,             # 页数（如果适用）
                    "word_count": int,        # 词数
                    "language": str,          # 语言
                }

        Raises:
            DocumentProcessingException: 提取失败
        """
        logger.info(f"提取文档内容: {file_path}", mime_type=mime_type)

        try:
            if mime_type == "application/pdf":
                return PDFExtractor.extract(file_path)

            elif mime_type in [
                "application/msword",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ]:
                return DOCXExtractor.extract(file_path)

            elif mime_type == "text/plain":
                return TextExtractor.extract(file_path)

            elif mime_type == "text/markdown":
                return MarkdownExtractor.extract(file_path)

            else:
                raise DocumentProcessingException(
                    error_code=ErrorCode.DOCUMENT_TYPE_UNSUPPORTED,
                    message=f"Unsupported document type: {mime_type}"
                )

        except Exception as e:
            logger.error(f"文档提取失败: {e}", file_path=file_path)
            raise DocumentProcessingException(
                error_code=ErrorCode.EXTRACTION_FAILED,
                message=f"Failed to extract content: {str(e)}"
            )


class PDFExtractor:
    """PDF 文档提取器"""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        """
        提取 PDF 内容

        Args:
            file_path: PDF 文件路径

        Returns:
            Dict: 提取结果
        """
        try:
            import PyPDF2

            text_content = []
            metadata = {}

            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)

                # 提取元数据
                if pdf_reader.metadata:
                    metadata = {
                        "title": pdf_reader.metadata.get('/Title', ''),
                        "author": pdf_reader.metadata.get('/Author', ''),
                        "subject": pdf_reader.metadata.get('/Subject', ''),
                        "creator": pdf_reader.metadata.get('/Creator', ''),
                    }

                # 提取文本
                num_pages = len(pdf_reader.pages)

                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()

                    if text.strip():
                        text_content.append(text)

            full_text = "\n\n".join(text_content)

            return {
                "text": full_text,
                "metadata": metadata,
                "pages": num_pages,
                "word_count": len(full_text.split()),
                "language": PDFExtractor._detect_language(full_text),
            }

        except Exception as e:
            logger.error(f"PDF 提取失败: {e}", file_path=file_path)
            raise

    @staticmethod
    def _detect_language(text: str) -> str:
        """
        检测文本语言

        Args:
            text: 文本内容

        Returns:
            str: 语言代码（如 "en", "zh"）
        """
        # 简单的语言检测（可以使用 langdetect 库）
        # 这里只做简单判断
        chinese_chars = sum(1 for c in text if '一' <= c <= '鿿')
        total_chars = len(text)

        if total_chars > 0 and chinese_chars / total_chars > 0.3:
            return "zh"
        else:
            return "en"


class DOCXExtractor:
    """DOCX 文档提取器"""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        """
        提取 DOCX 内容

        Args:
            file_path: DOCX 文件路径

        Returns:
            Dict: 提取结果
        """
        try:
            import docx

            doc = docx.Document(file_path)

            # 提取元数据
            metadata = {
                "title": doc.core_properties.title or "",
                "author": doc.core_properties.author or "",
                "subject": doc.core_properties.subject or "",
                "created": str(doc.core_properties.created) if doc.core_properties.created else "",
            }

            # 提取文本
            text_content = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)

            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_content.append(row_text)

            full_text = "\n".join(text_content)

            return {
                "text": full_text,
                "metadata": metadata,
                "pages": None,  # DOCX 没有页码概念
                "word_count": len(full_text.split()),
                "language": PDFExtractor._detect_language(full_text),
            }

        except Exception as e:
            logger.error(f"DOCX 提取失败: {e}", file_path=file_path)
            raise


class TextExtractor:
    """纯文本提取器"""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        """
        提取纯文本内容

        Args:
            file_path: 文本文件路径

        Returns:
            Dict: 提取结果
        """
        try:
            # 尝试多种编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']

            text = None
            used_encoding = None

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        text = f.read()
                    used_encoding = encoding
                    break
                except UnicodeDecodeError:
                    continue

            if text is None:
                raise ValueError("无法识别文件编码")

            return {
                "text": text,
                "metadata": {"encoding": used_encoding},
                "pages": None,
                "word_count": len(text.split()),
                "language": PDFExtractor._detect_language(text),
            }

        except Exception as e:
            logger.error(f"文本提取失败: {e}", file_path=file_path)
            raise


class MarkdownExtractor:
    """Markdown 文档提取器"""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        """
        提取 Markdown 内容

        Args:
            file_path: Markdown 文件路径

        Returns:
            Dict: 提取结果
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                markdown_text = f.read()

            # 提取标题作为元数据
            lines = markdown_text.split('\n')
            title = ""

            for line in lines:
                if line.startswith('# '):
                    title = line[2:].strip()
                    break

            # 转换为纯文本（移除 Markdown 标记）
            import re

            # 移除代码块
            text = re.sub(r'```[\s\S]*?```', '', markdown_text)

            # 移除行内代码
            text = re.sub(r'`[^`]+`', '', text)

            # 移除链接但保留文本
            text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)

            # 移除图片
            text = re.sub(r'!\[([^\]]*)\]\([^\)]+\)', '', text)

            # 移除 Markdown 标记
            text = re.sub(r'[#*_~]', '', text)

            return {
                "text": text,
                "metadata": {"title": title},
                "pages": None,
                "word_count": len(text.split()),
                "language": PDFExtractor._detect_language(text),
            }

        except Exception as e:
            logger.error(f"Markdown 提取失败: {e}", file_path=file_path)
            raise
